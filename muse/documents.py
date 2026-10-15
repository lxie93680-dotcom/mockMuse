"""Basic document parsing/creation in the authorized file workspace."""
import json
from .local_files import checked_path

MAX_BYTES=20*1024*1024
MAX_TEXT=100000

def process(action,path,content='',overwrite=False):
    target=checked_path(path); ext=target.suffix.lower()
    if action=='read':
        if target.stat().st_size>MAX_BYTES: raise ValueError('文档最多20 MiB')
        if ext=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(target)
            text='\n'.join((p.extract_text() or '') for p in reader.pages[:30])
            return {'text':text[:MAX_TEXT],'pages':len(reader.pages),'truncated':len(reader.pages)>30 or len(text)>MAX_TEXT,'note':'扫描PDF可能需要OCR，此工具没有OCR'}
        if ext=='.docx':
            from docx import Document
            doc=Document(target)
            text='\n'.join([p.text for p in doc.paragraphs]+['\t'.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
            return {'text':text[:MAX_TEXT],'truncated':len(text)>MAX_TEXT}
        if ext=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(target,read_only=True,data_only=False)
            try:
                sheets=[]
                for sheet in list(book)[:10]:
                    rows=[[str(v)[:2000] if v is not None else '' for v in row] for row in sheet.iter_rows(max_row=min(sheet.max_row,200),max_col=min(sheet.max_column,30),values_only=True)]
                    sheets.append({'name':sheet.title,'rows':rows,'truncated':sheet.max_row>200 or sheet.max_column>30})
                value=json.dumps({'sheets':sheets,'truncated':len(book.sheetnames)>10},ensure_ascii=False)
                if len(value)>MAX_TEXT: return {'text':value[:MAX_TEXT],'truncated':True}
                return json.loads(value)
            finally: book.close()
        raise ValueError('读取支持.pdf/.docx/.xlsx；不执行宏，不计算Excel公式')
    if action!='create': raise ValueError('action仅支持read/create')
    if len(content)>MAX_TEXT: raise ValueError('内容最多100000字符')
    if target.exists() and not overwrite: raise FileExistsError('目标已存在；明确授权后才能覆盖')
    if ext=='.docx':
        from docx import Document
        doc=Document()
        for line in content.split('\n'): doc.add_paragraph(line)
        doc.save(target)
    elif ext=='.pdf':
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        pdf=canvas.Canvas(str(target));pdf.setFont('STSong-Light',12);y=800
        for line in content.split('\n'):
            for i in range(0,max(1,len(line)),36):
                if y<50: pdf.showPage();pdf.setFont('STSong-Light',12);y=800
                pdf.drawString(40,y,line[i:i+36]);y-=18
        pdf.save()
    elif ext=='.xlsx':
        from openpyxl import Workbook
        rows=json.loads(content)
        if not isinstance(rows,list) or len(rows)>1000 or any(not isinstance(row,list) or len(row)>100 for row in rows):
            raise ValueError('Excel内容须为二维JSON数组，最多1000行100列')
        book=Workbook();sheet=book.active
        for row in rows:
            if any(v is not None and not isinstance(v,(str,int,float,bool)) for v in row): raise ValueError('单元格必须是文本、数字、布尔或null')
            sheet.append(row)
        book.save(target);book.close()
    else: raise ValueError('生成支持.docx/.pdf/.xlsx')
    return {'created':str(target),'bytes':target.stat().st_size}
